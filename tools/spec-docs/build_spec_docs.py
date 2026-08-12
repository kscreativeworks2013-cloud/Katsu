# -*- coding: utf-8 -*-
"""仕様書（docs/spec/*.md）を PDF / Word / PowerPoint に書き出す。

アプリ本体（src/）とは独立したツール。実行しても src/ には一切触れない。

使い方:
    pip install -r tools/spec-docs/requirements.txt
    python tools/spec-docs/build_spec_docs.py [出力ディレクトリ]

既定の出力先は dist/spec-docs/（.gitignore 済み）。

入力の書式:
    docs/spec/ の Markdown を **ファイル名順** に読み、1ファイル＝1セクションとして扱う。
    先頭の「# 見出し」がセクション見出し、それ以降が本文。
    本文は空行で段落に分かれる。PDF・Word・PowerPoint の3形式へ同じ本文を流し込むため、
    強調や表などの Markdown 記法は使わず、プレーンテキストで書く。
"""

from pathlib import Path
import sys

from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Pt as PPTPt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

REPO_ROOT = Path(__file__).resolve().parents[2]
SPEC_DIR = REPO_ROOT / "docs" / "spec"
DEFAULT_OUT = REPO_ROOT / "dist" / "spec-docs"

DOC_TITLE = "Luxury Beauty Visual Proposal OS"
DOC_SUBTITLE = "仕様書｜第1〜6章"
DOC_LEAD = "この資料は docs/spec/ の内容から自動生成した現時点の仕様書です。"

BASENAME = "Luxury_Beauty_Visual_Proposal_OS_仕様書"
# PowerPoint 1枚に流し込む最大行数。超えた分は「（続き）」スライドへ送る。
PPTX_LINES_PER_SLIDE = 14


def load_sections(spec_dir: Path) -> list[tuple[str, str]]:
    """docs/spec/*.md を (見出し, 本文) の並びとして読み込む。"""
    files = sorted(spec_dir.glob("*.md"))
    if not files:
        raise SystemExit(f"仕様書の Markdown が見つかりません: {spec_dir}")

    sections: list[tuple[str, str]] = []
    for path in files:
        text = path.read_text(encoding="utf-8").strip()
        first, _, body = text.partition("\n")
        if not first.startswith("# "):
            raise SystemExit(f"{path.name} の1行目は「# 見出し」である必要があります。")
        sections.append((first[2:].strip(), body.strip()))
    return sections


def build_pdf(sections: list[tuple[str, str]], path: Path) -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "title", parent=styles["Title"], fontName="HeiseiKakuGo-W5",
        fontSize=22, leading=28, alignment=TA_CENTER,
    )
    heading = ParagraphStyle(
        "heading", parent=styles["Heading1"], fontName="HeiseiKakuGo-W5",
        fontSize=15, leading=20, spaceAfter=8,
    )
    # wordWrap="CJK" がないと日本語が単語単位で折り返され、行末が版面からはみ出す。
    body = ParagraphStyle(
        "body", parent=styles["BodyText"], fontName="HeiseiKakuGo-W5",
        fontSize=9.5, leading=15, spaceAfter=10, wordWrap="CJK",
    )

    story = [
        Paragraph(DOC_TITLE, title),
        Spacer(1, 12),
        Paragraph(DOC_SUBTITLE, heading),
        Paragraph(DOC_LEAD, body),
    ]
    for head, text in sections:
        story += [PageBreak(), Paragraph(head, heading)]
        story += [Paragraph(p.replace("\n", "<br/>"), body) for p in text.split("\n\n")]

    SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40,
    ).build(story)


def build_docx(sections: list[tuple[str, str]], path: Path) -> None:
    doc = Document()
    # 東アジア用フォントを指定しないと、環境によって日本語が代替フォントに落ちる。
    normal = doc.styles["Normal"]
    normal.font.name = "Yu Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")

    doc.add_heading(DOC_TITLE, 0)
    doc.add_paragraph(DOC_SUBTITLE)
    doc.add_paragraph(DOC_LEAD)
    for head, text in sections:
        doc.add_heading(head, level=1)
        for para in text.split("\n\n"):
            doc.add_paragraph(para)
    doc.save(path)


def build_pptx(sections: list[tuple[str, str]], path: Path) -> None:
    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = DOC_TITLE
    cover.placeholders[1].text = DOC_SUBTITLE

    for head, text in sections:
        lines = [line for line in text.split("\n") if line.strip()]
        chunks = [
            lines[i : i + PPTX_LINES_PER_SLIDE]
            for i in range(0, len(lines), PPTX_LINES_PER_SLIDE)
        ] or [[""]]
        for index, chunk in enumerate(chunks):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = head if index == 0 else f"{head}（続き）"
            frame = slide.placeholders[1].text_frame
            frame.clear()
            frame.word_wrap = True
            for i, line in enumerate(chunk):
                para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                para.text = line
                para.font.size = PPTPt(13)
    prs.save(path)


def main() -> None:
    out = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_OUT
    out.mkdir(parents=True, exist_ok=True)

    sections = load_sections(SPEC_DIR)
    paths = {
        "PDF": out / f"{BASENAME}.pdf",
        "DOCX": out / f"{BASENAME}.docx",
        "PPTX": out / f"{BASENAME}.pptx",
    }
    build_pdf(sections, paths["PDF"])
    build_docx(sections, paths["DOCX"])
    build_pptx(sections, paths["PPTX"])

    for label, path in paths.items():
        print(f"{label}: {path}")


if __name__ == "__main__":
    main()
